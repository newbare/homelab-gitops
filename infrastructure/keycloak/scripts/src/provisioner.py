import csv
import logging
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path

log = logging.getLogger("provisioner")


def load_users(path: Path):
    suffix = path.suffix.lower()
    if suffix == ".csv":
        with path.open(newline="", encoding="utf-8") as f:
            return list(csv.DictReader(f))
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        headers = [str(h).strip() for h in rows[0]]
        return [dict(zip(headers, row)) for row in rows[1:] if any(row)]
    raise ValueError(f"Formato não suportado: {suffix}")


def process_user(client, user):
    username = user["username"].strip()
    role = user["role"].strip()
    try:
        user_id, action = client.upsert_user(user)
        client.assign_role(user_id, role)
        return username, action, role, None
    except Exception as e:
        return username, None, role, str(e)


def run(client, users, workers=8, dry_run=False):
    if dry_run:
        for u in users:
            log.info(f"[dry-run] {u['username']} -> {u['role']}")
        return {"created": 0, "updated": 0, "errors": 0}

    log.info("Autenticando no Keycloak...")
    client.authenticate()

    client.ensure_realm()
    client.ensure_client()
    client.ensure_roles(sorted({u["role"].strip() for u in users}))

    log.info(f"Processando {len(users)} usuários com {workers} workers...")
    stats = {"created": 0, "updated": 0, "errors": 0}

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = {pool.submit(process_user, client, u): u for u in users}
        for fut in as_completed(futures):
            username, action, role, error = fut.result()
            if error:
                stats["errors"] += 1
                log.error(f"[ERRO] {username} ({role}): {error}")
            else:
                stats[action] += 1
                log.info(f"[{action.upper()}] {username} -> {role}")

    return stats
