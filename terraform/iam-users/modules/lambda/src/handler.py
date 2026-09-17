"""
Lambda: process-users-xls

Lê uma planilha XLSX do S3 e cria:
  - Usuários IAM
  - Grupos IAM
  - Usuários SSO (Identity Store)
  - Grupos SSO (Identity Store)
  - Memberships

Gera arquivos catalog-info.yaml pro Backstage.
"""

import os
import io
import json
import boto3
import openpyxl
import yaml
from urllib.parse import unquote_plus

s3_client = boto3.client("s3")
iam_client = boto3.client("iam")
identitystore_client = boto3.client("identitystore")

IDENTITY_STORE_ID = os.environ["IDENTITY_STORE_ID"]
BUCKET_NAME = os.environ["BUCKET_NAME"]
ENVIRONMENT = os.environ.get("ENVIRONMENT", "dev")


def lambda_handler(event, context):
    """
    Entry point. Recebe evento S3.
    """
    print(f"Evento recebido: {json.dumps(event)}")

    # Extrai bucket + key do evento S3
    record = event["Records"][0]
    bucket = record["s3"]["bucket"]["name"]
    key = unquote_plus(record["s3"]["object"]["key"])

    print(f"Processando: s3://{bucket}/{key}")

    # Baixa o XLSX
    response = s3_client.get_object(Bucket=bucket, Key=key)
    xlsx_bytes = response["Body"].read()

    # Lê o XLSX
    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes))
    ws = wb.active

    # Cabeçalhos (linha 1)
    headers = [cell.value for cell in ws[1]]
    print(f"Cabeçalhos: {headers}")

    users = []
    groups = set()

    # Processa linhas (a partir da linha 2)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[0]:
            continue

        user_data = dict(zip(headers, row))
        users.append(user_data)

        # Coleta grupos (memberOf pode ter múltiplos separados por vírgula)
        member_of = user_data.get("memberOf", "")
        if member_of:
            for g in str(member_of).split(","):
                groups.add(g.strip())

    print(f"Total de usuários: {len(users)}")
    print(f"Total de grupos: {len(groups)}")

    # Cria os grupos no IAM e SSO
    for group_name in groups:
        create_iam_group(group_name)
        create_sso_group(group_name)

    # Cria os usuários
    for user_data in users:
        create_user(user_data)

    return {
        "statusCode": 200,
        "body": json.dumps({
            "users_created": len(users),
            "groups_created": len(groups),
        }),
    }


def create_iam_group(group_name):
    """Cria grupo no IAM (idempotente)."""
    try:
        iam_client.create_group(GroupName=group_name)
        print(f"✅ IAM group criado: {group_name}")
    except iam_client.exceptions.EntityAlreadyExistsException:
        print(f"⏭️  IAM group já existe: {group_name}")


def create_sso_group(group_name):
    """Cria grupo no Identity Store (idempotente)."""
    try:
        response = identitystore_client.create_group(
            IdentityStoreId=IDENTITY_STORE_ID,
            DisplayName=group_name,
        )
        print(f"✅ SSO group criado: {group_name} ({response['GroupId']})")
        return response["GroupId"]
    except Exception as e:
        if "ConflictException" in str(e):
            print(f"⏭️  SSO group já existe: {group_name}")
            return find_sso_group(group_name)
        raise


def find_sso_group(group_name):
    """Busca grupo no Identity Store pelo nome."""
    response = identitystore_client.list_groups(
        IdentityStoreId=IDENTITY_STORE_ID,
        Filters=[{"AttributePath": "DisplayName", "AttributeValue": group_name}],
    )
    groups = response.get("Groups", [])
    return groups[0]["GroupId"] if groups else None


def create_user(user_data):
    """Cria usuário no IAM + SSO + memberships + catalog-info.yaml."""
    username = str(user_data["name"]).strip()
    display_name = str(user_data["displayName"]).strip()
    email = str(user_data["email"]).strip()
    member_of = str(user_data.get("memberOf", "")).strip()

    # 1. IAM User
    try:
        iam_client.create_user(
            UserName=username,
            Tags=[
                {"Key": "Email", "Value": email},
                {"Key": "DisplayName", "Value": display_name},
                {"Key": "ManagedBy", "Value": "terraform-lambda"},
                {"Key": "Environment", "Value": ENVIRONMENT},
            ],
        )
        print(f"✅ IAM user criado: {username}")
    except iam_client.exceptions.EntityAlreadyExistsException:
        print(f"⏭️  IAM user já existe: {username}")

    # 2. Obtém o ARN do IAM user
    user_info = iam_client.get_user(UserName=username)
    user_arn = user_info["User"]["Arn"]

    # 3. SSO User (Identity Store)
    sso_user_id = None
    try:
        response = identitystore_client.create_user(
            IdentityStoreId=IDENTITY_STORE_ID,
            UserName=username,
            DisplayName=display_name,
            Emails=[{"Value": email, "Type": "work", "Primary": True}],
            Name={
                "GivenName": display_name.split(" ")[0],
                "FamilyName": " ".join(display_name.split(" ")[1:]) or "-",
            },
        )
        sso_user_id = response["UserId"]
        print(f"✅ SSO user criado: {username} ({sso_user_id})")
    except Exception as e:
        if "ConflictException" in str(e):
            print(f"⏭️  SSO user já existe: {username}")
        else:
            print(f"⚠️  Erro SSO user: {e}")

    # 4. Memberships (IAM + SSO)
    if member_of:
        groups_list = [g.strip() for g in member_of.split(",") if g.strip()]
        for group_name in groups_list:
            # IAM membership
            try:
                iam_client.add_user_to_group(
                    GroupName=group_name,
                    UserName=username,
                )
                print(f"✅ IAM membership: {username} → {group_name}")
            except Exception as e:
                print(f"⚠️  IAM membership: {e}")

            # SSO membership
            if sso_user_id:
                try:
                    sso_group_id = find_sso_group(group_name)
                    if sso_group_id:
                        identitystore_client.create_group_membership(
                            IdentityStoreId=IDENTITY_STORE_ID,
                            GroupId=sso_group_id,
                            MemberId={"UserId": sso_user_id},
                        )
                        print(f"✅ SSO membership: {username} → {group_name}")
                except Exception as e:
                    print(f"⚠️  SSO membership: {e}")

    # 5. Gera catalog-info.yaml
    catalog_yaml = generate_catalog_info(
        username=username,
        display_name=display_name,
        email=email,
        groups=[g.strip() for g in member_of.split(",") if g.strip()],
        user_arn=user_arn,
    )

    # 6. Upload do catalog-info.yaml no S3
    catalog_key = f"catalog/users/{username}.yaml"
    s3_client.put_object(
        Bucket=BUCKET_NAME,
        Key=catalog_key,
        Body=catalog_yaml,
        ContentType="application/x-yaml",
    )
    print(f"✅ catalog-info.yaml: s3://{BUCKET_NAME}/{catalog_key}")

    return user_arn


def generate_catalog_info(username, display_name, email, groups, user_arn):
    """Gera o YAML do Backstage (Kind: User)."""
    catalog = {
        "apiVersion": "backstage.io/v1alpha1",
        "kind": "User",
        "metadata": {
            "name": username,
            "annotations": {
                "aws-user-arn": user_arn,
            },
        },
        "spec": {
            "profile": {
                "displayName": display_name,
                "email": email,
            },
            "memberOf": groups,
        },
    }
    return yaml.dump(catalog, default_flow_style=False, sort_keys=False)
